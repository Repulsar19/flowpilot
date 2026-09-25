"""Extract text from PDF, DOCX and TXT files while preserving page/section hints."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# Matches numbered SOP headings such as "6. PROCEDURE", "6.3 Impact Assessment" or "6.3.2 The QA Analyst..."
_SECTION_RE = re.compile(r"^\s*(\d+(?:\.\d+){0,3})\.?\s+([A-Z][^\n]{2,120}?)\s*$")


@dataclass
class DocumentSection:
    label: str
    title: str
    page: int | None
    start_char: int


@dataclass
class ParsedDocument:
    text: str
    page_count: int | None
    sections: list[DocumentSection] = field(default_factory=list)

    @property
    def preview(self) -> str:
        return self.text[:600].strip()


class UnsupportedDocumentError(ValueError):
    pass


class DocumentParser:
    """Stateless parser. Add new formats by extending `_extract_*` and SUPPORTED_EXTENSIONS."""

    def parse(self, path: Path) -> ParsedDocument:
        ext = path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentError(f"Unsupported file type: {ext}")

        if ext == ".pdf":
            pages = self._extract_pdf(path)
        elif ext == ".docx":
            pages = [self._extract_docx(path)]
        else:
            pages = [path.read_text(encoding="utf-8", errors="replace")]

        text, sections = self._join_pages(pages)
        return ParsedDocument(text=text, page_count=len(pages) if ext == ".pdf" else None, sections=sections)

    def _extract_pdf(self, path: Path) -> list[str]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [(page.extract_text() or "") for page in reader.pages]

    def _extract_docx(self, path: Path) -> str:
        import docx

        document = docx.Document(str(path))
        parts: list[str] = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)

    def _join_pages(self, pages: list[str]) -> tuple[str, list[DocumentSection]]:
        chunks: list[str] = []
        sections: list[DocumentSection] = []
        offset = 0
        for idx, page_text in enumerate(pages, start=1):
            for line in page_text.splitlines():
                match = _SECTION_RE.match(line)
                if match:
                    sections.append(
                        DocumentSection(
                            label=match.group(1),
                            title=match.group(2).strip(),
                            page=idx if len(pages) > 1 else None,
                            start_char=offset,
                        )
                    )
                offset += len(line) + 1
            chunks.append(page_text)
            offset += 1
        return "\n".join(chunks), sections
